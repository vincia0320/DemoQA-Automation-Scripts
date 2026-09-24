from pathlib import Path
from openpyxl import load_workbook
from playwright.sync_api import expect, sync_playwright

BASE = Path(__file__).resolve().parent
PROJECT_ROOT = BASE.parent
INPUT_FILE = PROJECT_ROOT / "Forms Uploader.xlsx"

def get_input_file() -> Path:
    if not INPUT_FILE.is_file():
        raise FileNotFoundError(f"Could not find '{INPUT_FILE}'")
    return INPUT_FILE


def first_value(row, *names):
    values = {str(k).strip().lower(): v for k, v in row.items() if k is not None}
    for name in names:
        value = values.get(name.lower())
        if value not in (None, ""):
            return str(value)
    return ""


def main():
    input_file = get_input_file()
    workbook = load_workbook(input_file, data_only=True)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    row = dict(zip(headers, next(sheet.iter_rows(min_row=2, max_row=2, values_only=True))))

    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=False)
        page = browser.new_page()

        # Open the form in the same browser instance
        page.goto("https://demoqa.com/automation-practice-form", wait_until="domcontentloaded")

        fields = {
            "firstName": ("firstName", "first name"),
            "lastName": ("lastName", "last name"),
            "userEmail": ("email", "user email"),
            "userNumber": ("mobile", "mobile number", "phone"),
            "currentAddress": ("address", "current address"),
        }
        for element_id, names in fields.items():
            value = first_value(row, *names)
            if value:
                page.locator(f"#{element_id}").fill(value)

        gender = first_value(row, "gender")
        if gender:
            page.locator(f"label:text-is('{gender}')").click()

        subjects = first_value(row, "subjects", "subject")
        if subjects:
            subject_input = page.locator("#subjectsInput")
            subject_input.fill(subjects)
            subject_input.press("Enter")

        for element_id in ("state", "city"):
            value = first_value(row, element_id)
            if value:
                page.locator(f"#{element_id}").click()
                page.get_by_text(value, exact=True).click()

        page.locator("#submit").click()
        expect(page.locator("#example-modal-sizes-title-lg")).to_be_visible()
        print("Passed")
        browser.close()


if __name__ == "__main__":
    main()
